from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\thirs\Documents\Calculo MS Tesla")
INPUT = ROOT / "inputs" / "data (33).xlsx"
OUTPUT_DIR = ROOT / "outputs" / "market_study"
OUTPUT = OUTPUT_DIR / "Tesla Competition Study - Passenger Vehicles Under 25MM - 2026.xlsx"

YEAR = 2026
PRICE_MAX = 25_000_000
EXCLUDED_MODELS = {"RIFTER", "EXPERT"}

STARTING_PRICE = {
    ("SUBARU", "IMPREZA SPORT NEW GENERATION"): 19_900_000,
    ("CITROEN", "NEW C4"): 22_900_000,
    ("PEUGEOT", "308"): 19_900_000,
    ("TOYOTA", "COROLLA"): 21_990_000,
}


def normalize_text(value):
    if pd.isna(value):
        return ""
    return str(value).strip().upper()


def weighted_average(group):
    demand = group["demand"].sum()
    if demand == 0:
        return group["price"].mean()
    return (group["price"] * group["demand"]).sum() / demand


def build_table():
    raw = pd.read_excel(INPUT, sheet_name="Export")
    data = pd.DataFrame(
        {
            "year": pd.to_numeric(raw.iloc[:, 0], errors="coerce"),
            "brand": raw.iloc[:, 2].map(normalize_text),
            "model": raw.iloc[:, 3].map(normalize_text),
            "price": pd.to_numeric(raw.iloc[:, 6], errors="coerce"),
            "demand": pd.to_numeric(raw.iloc[:, 7], errors="coerce").fillna(0),
            "segment": raw.iloc[:, 10].astype(str),
        }
    )

    passenger = data[
        (data["year"] == YEAR)
        & data["segment"].str.contains("Pasaj", case=False, na=False)
        & data["price"].gt(0)
        & ~data["model"].isin(EXCLUDED_MODELS)
    ].copy()

    passenger["weighted_price"] = passenger["price"] * passenger["demand"]
    grouped = passenger.groupby(["brand", "model"], as_index=False).agg(
        total_weighted_price=("weighted_price", "sum"),
        **{"Annual Demand 2026": ("demand", "sum")},
    )
    grouped["Average Anac List Price"] = (
        grouped["total_weighted_price"] / grouped["Annual Demand 2026"]
    )
    grouped = grouped.drop(columns=["total_weighted_price"])

    grouped = grouped[grouped["Annual Demand 2026"] > 0].copy()
    grouped = grouped[grouped["Average Anac List Price"] <= PRICE_MAX].copy()
    grouped["Starting price"] = grouped.apply(
        lambda row: STARTING_PRICE.get((row["brand"], row["model"])), axis=1
    )
    grouped = grouped.sort_values("Annual Demand 2026", ascending=False).reset_index(drop=True)
    return grouped


def write_workbook(grouped):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Passenger Vehicles 2026"

    headers = [
        "Ranking",
        "Brand",
        "Model",
        "Average Anac List Price",
        "Starting price",
        "Annual Demand 2026",
        "MS 2026",
    ]

    market_total = int(grouped["Annual Demand 2026"].sum())
    models_considered = len(grouped)
    top = grouped.head(10).copy()
    rest = grouped.iloc[10:].copy()

    rows = []
    for idx, row in top.iterrows():
        rows.append(
            [
                idx + 1,
                row["brand"],
                row["model"],
                float(row["Average Anac List Price"]),
                row["Starting price"],
                int(row["Annual Demand 2026"]),
                None,
            ]
        )

    if not rest.empty:
        rest_demand = int(rest["Annual Demand 2026"].sum())
        rest_price = float((rest["Average Anac List Price"] * rest["Annual Demand 2026"]).sum() / rest_demand)
        rows.append([None, "Rest", "Resting Models", rest_price, None, rest_demand, None])

    max_col = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"] = "Tesla Competition Study | Passenger Vehicles | 2026"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
    ws["A3"] = "Passenger Vehicles | bracket $0 to $25,000,000 CLP | 2026"
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max_col)
    ws["A4"] = f"2026: {market_total:,} Annual Demand | Models considered: {models_considered}"

    header_row = 6
    for col, header in enumerate(headers, start=1):
        ws.cell(header_row, col).value = header

    first_data_row = header_row + 1
    for ridx, row in enumerate(rows, start=first_data_row):
        for cidx, value in enumerate(row, start=1):
            ws.cell(ridx, cidx).value = value

    last_data_row = first_data_row + len(rows) - 1
    for ridx in range(first_data_row, last_data_row + 1):
        ws.cell(ridx, 7).value = f"=F{ridx}/SUM($F${first_data_row}:$F${last_data_row})"

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ["A1"]:
        ws[cell].font = Font(bold=True, color="FFFFFF", size=14)
        ws[cell].fill = title_fill
        ws[cell].alignment = Alignment(horizontal="center")

    for cell in ["A3", "A4"]:
        ws[cell].font = Font(bold=True, color="1F1F1F")
        ws[cell].fill = section_fill
        ws[cell].alignment = Alignment(horizontal="center")

    for col in range(1, max_col + 1):
        cell = ws.cell(header_row, col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row in range(first_data_row, last_data_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if col in [1, 6, 7] else "left")
        ws.cell(row, 4).number_format = '"$"#,##0'
        ws.cell(row, 5).number_format = '"$"#,##0'
        ws.cell(row, 6).number_format = "#,##0"
        ws.cell(row, 7).number_format = "0.0%"

    for row in [1, 3, 4]:
        ws.row_dimensions[row].height = 24
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False
    widths = [10, 18, 32, 24, 18, 20, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    wb.save(OUTPUT)
    return market_total, models_considered


if __name__ == "__main__":
    table = build_table()
    total, count = write_workbook(table)
    print(OUTPUT)
    print("models_considered", count, "market_total", total)
    print(table.head(12).to_string(index=False))
